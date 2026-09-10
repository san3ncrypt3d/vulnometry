"""Registry, reports, providers and the analyst loop."""

import asyncio
import json

from vulnometry.analyst import run_analysis
from vulnometry.providers import parse_spec
from vulnometry.providers.base import ChatResponse, Message, ProviderError, ToolCall, parse_arguments
from vulnometry.registry import ACTIONS, invoke
from vulnometry.report import build_dashboard, build_workbook, to_csv, to_json, to_markdown
from vulnometry.schema import (
    Assessment,
    ConfirmedExploitation,
    ExploitProbability,
    ExposureMeasure,
    Weakness,
)
from vulnometry.surfaces.schemas import DIALECTS


def test_every_action_schema_is_well_formed():
    for action in ACTIONS:
        schema = action.json_schema()
        assert schema["type"] == "object"
        assert isinstance(schema["properties"], dict)
        for required in schema.get("required", []):
            assert required in schema["properties"], f"{action.name}: {required} undeclared"
        assert len(action.description) > 40, f"{action.name} needs a usable description"


def test_every_dialect_exports_all_actions():
    count = len(ACTIONS)
    assert len(DIALECTS["openai"]()) == count
    assert len(DIALECTS["anthropic"]()) == count
    assert len(DIALECTS["bedrock"]()["tools"]) == count
    assert len(DIALECTS["gemini"]()[0]["function_declarations"]) == count


def test_gemini_schema_drops_unsupported_keywords():
    blob = json.dumps(DIALECTS["gemini"]())
    assert "additionalProperties" not in blob, "Gemini rejects this keyword"


def test_unknown_action_returns_data_not_an_exception():
    result = asyncio.run(invoke("no_such_action", {}))
    assert "error" in result
    assert "measure_exposure" in result["available"]


def test_bad_arguments_return_data_not_an_exception():
    result = asyncio.run(invoke("measure_exposure", {"cve_id": "banana"}))
    assert "error" in result and "banana" in result["error"]


def test_unexpected_arguments_are_ignored_with_a_warning():
    result = asyncio.run(invoke("harvest_cve_ids", {"text": "CVE-2021-44228", "nonsense": 1}))
    assert result["cve_ids"] == ["CVE-2021-44228"]
    assert "nonsense" in result["_warning"]


def test_offline_action_needs_no_network():
    result = asyncio.run(invoke("harvest_cve_ids", {"text": "saw CVE-2021-44228 in prod"}))
    assert result == {"count": 1, "cve_ids": ["CVE-2021-44228"]}


def test_model_spec_parsing_survives_colons_in_model_ids():
    assert parse_spec("bedrock:us.anthropic.claude-sonnet-4-20250514-v1:0") == (
        "bedrock", "us.anthropic.claude-sonnet-4-20250514-v1:0")
    assert parse_spec("ollama:qwen3:8b") == ("ollama", "qwen3:8b")
    assert parse_spec("azure:my-deployment") == ("azure", "my-deployment")
    assert parse_spec("llama3.1")[0] == "ollama", "a bare name defaults to Ollama"


def test_tool_argument_parsing_is_forgiving():
    assert parse_arguments('{"cve_id": "CVE-2021-44228"}') == {"cve_id": "CVE-2021-44228"}
    assert parse_arguments('```json\n{"a": 1}\n```') == {"a": 1}
    assert parse_arguments("total garbage") == {}
    assert parse_arguments({"already": "a dict"}) == {"already": "a dict"}


def test_each_provider_builds_its_own_wire_format():
    from vulnometry.providers import anthropic, bedrock, gemini, openai_compat

    convo = [
        Message(role="system", content="brief"),
        Message(role="user", content="is it urgent?"),
        Message(role="assistant", tool_calls=[ToolCall(id="c1", name="measure_exposure", arguments={"cve_id": "CVE-1"})]),
        Message(role="tool", name="measure_exposure", tool_call_id="c1", content='{"index": 900}'),
    ]

    openai_wire = openai_compat._to_wire(convo)
    assert openai_wire[2]["tool_calls"][0]["function"]["name"] == "measure_exposure"
    assert openai_wire[3] == {"role": "tool", "tool_call_id": "c1", "content": '{"index": 900}'}

    system, anthropic_wire = anthropic._to_wire(convo)
    assert system == "brief"
    assert anthropic_wire[1]["content"][0]["type"] == "tool_use"
    assert anthropic_wire[2]["content"][0]["type"] == "tool_result"

    bedrock_system, bedrock_wire = bedrock._to_wire(convo)
    assert bedrock_system == [{"text": "brief"}]
    assert "toolUse" in bedrock_wire[1]["content"][0]
    assert "toolResult" in bedrock_wire[2]["content"][0]

    instruction, contents = gemini._to_wire(convo)
    assert instruction["parts"][0]["text"] == "brief"
    assert "functionCall" in contents[1]["parts"][0]
    assert "functionResponse" in contents[2]["parts"][0]


class Scripted:
    name = "fake"
    model = "scripted"

    def __init__(self, script):
        self.script = list(script)
        self.received = []

    async def chat(self, messages, tools=None, **kwargs):
        self.received.append((list(messages), tools))
        return self.script.pop(0)

    async def check(self):
        return {"ok": True}

    async def aclose(self):
        pass


def test_analyst_calls_an_action_then_answers():
    provider = Scripted([
        ChatResponse(tool_calls=[ToolCall(id="c1", name="harvest_cve_ids",
                                          arguments={"text": "we saw CVE-2021-44228"})]),
        ChatResponse(text="One CVE: CVE-2021-44228."),
    ])
    run = asyncio.run(run_analysis("what is in this?", provider=provider))

    assert run.calls == 1
    assert run.actions_used() == ["harvest_cve_ids"]
    tool_messages = [m for m in provider.received[-1][0] if m.role == "tool"]
    assert tool_messages[0].tool_call_id == "c1"
    assert "CVE-2021-44228" in tool_messages[0].content


def test_analyst_survives_a_hallucinated_action():
    provider = Scripted([
        ChatResponse(tool_calls=[ToolCall(id="c1", name="does_not_exist", arguments={})]),
        ChatResponse(text="That action is unavailable."),
    ])
    run = asyncio.run(run_analysis("go", provider=provider))
    assert run.answer.startswith("That action is unavailable")
    assert "unknown action" in [m for m in provider.received[-1][0] if m.role == "tool"][0].content


def test_analyst_stops_looping_and_withholds_actions_on_the_last_call():
    loop = ChatResponse(tool_calls=[ToolCall(id="c", name="harvest_cve_ids", arguments={"text": "x"})])
    provider = Scripted([loop] * 3 + [ChatResponse(text="Final answer.")])
    run = asyncio.run(run_analysis("go", provider=provider, max_turns=3))
    assert run.turns == 3
    assert run.answer == "Final answer."
    assert provider.received[-1][1] is None


def test_provider_failure_becomes_an_answer_not_a_crash():
    class Broken(Scripted):
        async def chat(self, messages, tools=None, **kwargs):
            raise ProviderError("ollama", "connection refused", "run `ollama serve`")

    run = asyncio.run(run_analysis("go", provider=Broken([])))
    assert "Inference failed" in run.answer and "ollama serve" in run.answer


def _sample() -> list[Assessment]:
    return [
        Assessment(cve_id="CVE-2021-44228", asset="checkout-api", owner="pay@x.com",
                   business_unit="Commerce", due_by="2026-09-04", sla_days=2,
                   directive="Act now.", fixes=["log4j-core 2.15.0"],
                   weakness=Weakness(cve_id="CVE-2021-44228", summary="RCE", resolved=True),
                   exposure=ExposureMeasure(index=928.9, verdict="Contain", threat=1.0,
                                            reachability=0.9, consequence=1.0,
                                            threat_basis=["CISA KEV lists this as exploited"])),
        Assessment(cve_id="CVE-2023-40000", asset="ml-sandbox", directive="No action.",
                   exposure=ExposureMeasure(index=0.0, verdict="Accept", threat=0.26,
                                            reachability=0.0, consequence=0.2,
                                            collapsed_by="not-deployed")),
    ]


def test_workbook_has_the_expected_sheets_and_survives_a_round_trip(tmp_path):
    from openpyxl import load_workbook

    path = build_workbook(_sample(), tmp_path / "out.xlsx",
                          {"assessed": 2, "actionable": 1, "suppressed_by_context": 1, "unattributed": 0})
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Findings", "Action Plan", "Accepted", "By Owner", "Method"]

    findings = workbook["Findings"]
    assert findings.freeze_panes == "B5"
    assert findings.auto_filter.ref
    assert findings.cell(row=5, column=1).value == "CVE-2021-44228"
    assert findings.cell(row=5, column=3).value == "Contain"
    assert "Rationale" in [c.value for c in findings[4]]
    assert workbook["Action Plan"].cell(row=5, column=2).value == "CVE-2021-44228"
    assert workbook["Action Plan"].cell(row=6, column=2).value is None


def test_accepted_sheet_carries_rationale_and_signoff_columns(tmp_path):
    from openpyxl import load_workbook

    accepted = Assessment(
        cve_id="CVE-2024-9999", asset="reporting-svc", tier=3,
        internet_exposed=False, data_classification="internal",
        weakness=Weakness(cve_id="CVE-2024-9999", summary="XSS", resolved=True),
        probability=ExploitProbability(cve_id="CVE-2024-9999", probability=0.004,
                                       as_of="2026-09-08", resolved=True),
        exposure=ExposureMeasure(index=96, verdict="Accept", threat=0.13, reachability=0.7,
                                 consequence=0.5, confidence="high",
                                 threat_basis=["EPSS puts exploitation within 30 days at 0.40%"],
                                 reachability_basis=["Attack vector is network"],
                                 consequence_basis=["CVSS impact C:LOW I:LOW A:NONE"]),
        measured_at="2026-09-09T12:00:00+00:00",
    )
    path = build_workbook([accepted], tmp_path / "out.xlsx", None)
    sheet = load_workbook(path)["Accepted"]
    header = [c.value for c in sheet[4]]
    for col in ("Rationale", "Internet", "EPSS %", "EPSS date", "Reopen triggers", "Decision", "Model"):
        assert col in header, col

    row = {header[i]: sheet.cell(row=5, column=i + 1).value for i in range(len(header))}
    assert row["CVE"] == "CVE-2024-9999"
    assert row["Internet"] == "Internal"
    assert row["EPSS date"] == "2026-09-08"
    assert row["Model"] == "bei-1.0"
    assert "little sign anyone will try" in row["Rationale"]
    assert "EPSS 0.40% (as of 2026-09-08)" in row["Rationale"]
    assert row["Decision"] == "Accept risk / suppress in scanner"
    assert "EPSS >= 0.10" in row["Reopen triggers"]
    assert "CISA KEV" in row["Reopen triggers"]


def test_rationale_leads_with_the_driving_factor():
    urgent = Assessment(
        cve_id="CVE-1", internet_exposed=True, tier=1,
        exploitation=ConfirmedExploitation(cve_id="CVE-1", confirmed=True),
        exposure=ExposureMeasure(index=800, verdict="Contain", threat=1.0, reachability=0.9,
                                 consequence=0.9, confidence="high",
                                 threat_basis=["CISA KEV lists this as exploited in the wild"]),
    )
    assert urgent.rationale().startswith("Contain: exploitation is likely or under way")
    assert "internet-facing" in urgent.rationale()
    assert "in CISA KEV" in urgent.rationale()


def test_dashboard_is_self_contained(tmp_path):
    path = build_dashboard(_sample(), tmp_path / "out.html",
                           {"assessed": 2, "actionable": 1, "suppressed_by_context": 1,
                            "unattributed": 0, "confirmed_exploited": ["CVE-2021-44228"],
                            "load_by_owner": {"pay@x.com": 1},
                            "exposure_by_business_unit": {"Commerce": 928.9}})
    document = path.read_text()
    assert "<script" not in document, "no JavaScript: it must open from file:// and survive email"
    assert "cdn" not in document.lower()
    assert document.count("https://") == document.count("http://www.w3.org")
    assert "CVE-2021-44228" in document and "Contain" in document
    assert "<svg" in document


def test_html_escaping_in_the_dashboard(tmp_path):
    items = _sample()
    items[0].asset = "<script>alert(1)</script>"
    document = build_dashboard(items, tmp_path / "x.html", {}).read_text()
    assert "<script>alert(1)</script>" not in document
    assert "&lt;script&gt;" in document


def test_text_formats_render():
    items = _sample()
    assert json.loads(to_json(items, {"assessed": 2}))["summary"]["assessed"] == 2
    markdown = to_markdown(items, {"assessed": 2, "actionable": 1, "deferrable": 1,
                                   "contain_now": ["CVE-2021-44228"], "suppressed_by_context": 1})
    assert "| CVE-2021-44228 | 928.9 | Contain" in markdown
    assert "## CVE-2021-44228" in markdown
    csv_text = to_csv(items)
    assert csv_text.splitlines()[0].startswith("cve,bei,verdict")
    assert "CVE-2021-44228,928.9,Contain" in csv_text


def test_the_package_version_matches_pyproject():
    """__version__ reaches users through `vulnometry --version`, the MCP
    serverInfo handshake and the OpenAPI document, and it is a second copy of a
    number that lives in pyproject.toml. The 0.2.0 release bumped one and not
    the other, so every one of those three reported 0.1.0.
    """
    import re
    from pathlib import Path

    import vulnometry

    pyproject = Path(__file__).resolve().parent.parent / "pyproject.toml"
    declared = re.search(r'^version = "([^"]+)"', pyproject.read_text(), re.M)
    assert declared, "pyproject.toml has no top-level version"
    assert vulnometry.__version__ == declared.group(1)
