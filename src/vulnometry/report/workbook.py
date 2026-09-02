"""Annotated Excel output: Findings, Action Plan, By Owner, Method."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..schema import Assessment

VERDICT_FILL = {
    "Contain": PatternFill("solid", fgColor="C0392B"),
    "Remediate": PatternFill("solid", fgColor="E67E22"),
    "Schedule": PatternFill("solid", fgColor="F1C40F"),
    "Accept": PatternFill("solid", fgColor="BDC3C7"),
}
VERDICT_FONT = {
    "Contain": Font(color="FFFFFF", bold=True),
    "Remediate": Font(color="FFFFFF", bold=True),
    "Schedule": Font(color="4A3B00"),
    "Accept": Font(color="4A4A4A"),
}

HEADER_FILL = PatternFill("solid", fgColor="1F3A53")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
THIN = Side(style="thin", color="D5D8DC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FINDINGS_COLUMNS = [
    ("CVE", 16), ("BEI", 8), ("Verdict", 12), ("Asset", 22), ("Owner", 26),
    ("Business unit", 18), ("Environment", 13), ("Due by", 12), ("SLA days", 9),
    ("Threat", 9), ("Reach", 9), ("Consequence", 12),
    ("CVSS", 7), ("EPSS %", 9), ("KEV", 6), ("Ransomware", 11),
    ("Component", 22), ("Fixed in", 28), ("Confidence", 11),
    ("Primary driver", 52), ("Directive", 60),
]


def _style_header(sheet, row_index: int, count: int) -> None:
    for column in range(1, count + 1):
        cell = sheet.cell(row=row_index, column=column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER


def _widths(sheet, columns) -> None:
    for index, (_, width) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _primary_driver(item: Assessment) -> str:
    if item.exposure.collapsed_by == "not-deployed":
        return "Suppressed: component not deployed on this asset"
    for basis in (item.exposure.threat_basis, item.exposure.reachability_basis):
        if basis:
            return basis[0]
    return ""


def _findings_sheet(workbook: Workbook, items: list[Assessment]) -> None:
    sheet = workbook.active
    sheet.title = "Findings"

    sheet.cell(row=1, column=1, value="Exposure assessment").font = TITLE_FONT
    sheet.cell(row=2, column=1, value=f"Generated {date.today().isoformat()} · Business Exposure Index 0-1000 · higher means act sooner")
    sheet.cell(row=2, column=1).font = Font(color="7F8C8D", size=10)

    header_row = 4
    for index, (name, _) in enumerate(FINDINGS_COLUMNS, start=1):
        sheet.cell(row=header_row, column=index, value=name)
    _style_header(sheet, header_row, len(FINDINGS_COLUMNS))

    for offset, item in enumerate(items):
        row = header_row + 1 + offset
        severity = item.weakness.primary_severity()
        values = [
            item.cve_id,
            item.exposure.index,
            item.exposure.verdict,
            item.asset or "",
            item.owner or "",
            item.business_unit or "",
            item.environment or "",
            item.due_by or "",
            item.sla_days or "",
            item.exposure.threat,
            item.exposure.reachability,
            item.exposure.consequence,
            severity.base_score if severity and severity.base_score is not None else "",
            round(item.probability.probability * 100, 2) if item.probability.probability is not None else "",
            "Yes" if item.exploitation.confirmed else "",
            "Yes" if item.exploitation.ransomware_linked else "",
            (item.bulletins[0].remedies[0].component if item.bulletins and item.bulletins[0].remedies else ""),
            ", ".join(item.fixes[:3]),
            item.exposure.confidence,
            _primary_driver(item),
            item.directive,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=column >= 20)

        verdict_cell = sheet.cell(row=row, column=3)
        verdict_cell.fill = VERDICT_FILL.get(item.exposure.verdict, VERDICT_FILL["Accept"])
        verdict_cell.font = VERDICT_FONT.get(item.exposure.verdict, VERDICT_FONT["Accept"])
        verdict_cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.cell(row=row, column=2).font = Font(bold=True)

    _widths(sheet, FINDINGS_COLUMNS)
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=2)
    if items:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(FINDINGS_COLUMNS))}{header_row + len(items)}"


def _action_sheet(workbook: Workbook, items: list[Assessment]) -> None:
    actionable = [i for i in items if i.exposure.verdict in ("Contain", "Remediate")]
    sheet = workbook.create_sheet("Action Plan")

    sheet.cell(row=1, column=1, value="What has to happen, and by when").font = TITLE_FONT
    sheet.cell(row=2, column=1, value=f"{len(actionable)} of {len(items)} findings need action. Sorted by due date.")
    sheet.cell(row=2, column=1).font = Font(color="7F8C8D", size=10)

    columns = [("Due by", 12), ("CVE", 16), ("Verdict", 12), ("Asset", 24), ("Owner", 28), ("Action", 70)]
    header_row = 4
    for index, (name, _) in enumerate(columns, start=1):
        sheet.cell(row=header_row, column=index, value=name)
    _style_header(sheet, header_row, len(columns))

    ordered = sorted(actionable, key=lambda i: (i.due_by or "9999-12-31", -i.exposure.index))
    for offset, item in enumerate(ordered):
        row = header_row + 1 + offset
        for column, value in enumerate(
            [item.due_by or "unscheduled", item.cve_id, item.exposure.verdict,
             item.asset or "unattributed", item.owner or "unassigned", item.directive],
            start=1,
        ):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=column == 6)
        cell = sheet.cell(row=row, column=3)
        cell.fill = VERDICT_FILL.get(item.exposure.verdict, VERDICT_FILL["Accept"])
        cell.font = VERDICT_FONT.get(item.exposure.verdict, VERDICT_FONT["Accept"])

    _widths(sheet, columns)
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)


def _owner_sheet(workbook: Workbook, items: list[Assessment]) -> None:
    sheet = workbook.create_sheet("By Owner")
    sheet.cell(row=1, column=1, value="Workload by owner").font = TITLE_FONT
    sheet.cell(row=2, column=1, value="Who is carrying the queue, and how urgent their share is.").font = Font(color="7F8C8D", size=10)

    tally: dict[str, dict] = {}
    for item in items:
        owner = item.owner or "unassigned"
        bucket = tally.setdefault(owner, {"Contain": 0, "Remediate": 0, "Schedule": 0, "Accept": 0, "exposure": 0.0, "soonest": ""})
        bucket[item.exposure.verdict] = bucket.get(item.exposure.verdict, 0) + 1
        bucket["exposure"] += item.exposure.index
        if item.due_by and (not bucket["soonest"] or item.due_by < bucket["soonest"]):
            bucket["soonest"] = item.due_by

    columns = [("Owner", 32), ("Contain", 10), ("Remediate", 11), ("Schedule", 10), ("Accept", 9), ("Total exposure", 15), ("Earliest due", 13)]
    header_row = 4
    for index, (name, _) in enumerate(columns, start=1):
        sheet.cell(row=header_row, column=index, value=name)
    _style_header(sheet, header_row, len(columns))

    ordered = sorted(tally.items(), key=lambda kv: (-kv[1]["Contain"], -kv[1]["exposure"]))
    for offset, (owner, bucket) in enumerate(ordered):
        row = header_row + 1 + offset
        for column, value in enumerate(
            [owner, bucket["Contain"], bucket["Remediate"], bucket["Schedule"],
             bucket["Accept"], round(bucket["exposure"], 1), bucket["soonest"] or ""],
            start=1,
        ):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.border = BORDER
        if bucket["Contain"]:
            sheet.cell(row=row, column=2).fill = VERDICT_FILL["Contain"]
            sheet.cell(row=row, column=2).font = VERDICT_FONT["Contain"]

    _widths(sheet, columns)


def _method_sheet(workbook: Workbook, items: list[Assessment], summary: dict | None) -> None:
    sheet = workbook.create_sheet("Method")
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 96

    rows = [
        ("Business Exposure Index", "BEI = 1000 x Threat^0.8 x Reachability^0.7 x Consequence^0.6"),
        ("", "Multiplicative, so any factor can veto. If nobody can reach it, exposure collapses regardless of CVSS."),
        ("Threat", "Will anyone try? CISA KEV listing = 1.0. Otherwise the stronger of EPSS (concave) and exploit maturity."),
        ("Reachability", "Can they get to it here? CVSS attack vector, complexity, privileges and user interaction, "
                         "adjusted by internet exposure, deployment status and compensating controls."),
        ("Consequence", "What does it cost us? CVSS impact sub-metrics, scaled by asset tier, data classification, "
                        "regulatory scope and environment."),
        ("", ""),
        ("Verdict thresholds", "Contain >= 700 · Remediate >= 400 · Schedule >= 150 · Accept < 150"),
        ("Contain", "Act now, outside the normal change process."),
        ("Remediate", "Fix within this sprint, ahead of the routine cycle."),
        ("Schedule", "Queue for the next planned maintenance window."),
        ("Accept", "No action warranted. Record the decision and move on."),
        ("", ""),
        ("Due dates", "Derived from the verdict and the asset: tier 1 halves the window, regulated systems take 60% "
                      "of it, production takes 80%. Not a guess, but a policy you can edit in the inventory."),
        ("Confidence", "high = CVE record, EPSS and asset context all present. low = at least two are missing; "
                       "treat the number as provisional."),
        ("", ""),
        ("Data sources", "NIST NVD, CVE Program (fallback), FIRST EPSS, CISA KEV, OSV.dev, GitHub Security Advisories."),
        ("Caveat", "Exploit-artifact discovery is a heuristic GitHub search. A match means public code claims to "
                   "exploit the CVE, not that it works."),
        ("Caveat", "Asset context comes from your inventory. Nothing here can verify that the vulnerable code path "
                   "is actually reachable in your deployment."),
    ]
    if summary:
        rows += [
            ("", ""),
            ("This assessment", f"{summary['assessed']} findings · {summary['actionable']} actionable · "
                                f"{summary['suppressed_by_context']} suppressed by business context · "
                                f"{summary['unattributed']} not matched to a known asset"),
        ]

    sheet.cell(row=1, column=1, value="How these numbers were produced").font = TITLE_FONT
    for offset, (label, text) in enumerate(rows, start=3):
        sheet.cell(row=offset, column=1, value=label).font = Font(bold=bool(label))
        cell = sheet.cell(row=offset, column=2, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_workbook(items: list[Assessment], path: str | Path, summary: dict | None = None) -> Path:
    path = Path(path)
    workbook = Workbook()
    _findings_sheet(workbook, items)
    _action_sheet(workbook, items)
    _owner_sheet(workbook, items)
    _method_sheet(workbook, items, summary)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
